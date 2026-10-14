import io
import socket
import zipfile
import pytest
from docx import Document as WordDocument
from pydantic import ValidationError
from docaudit import audit, parse_document, load_config, AuditConfig, DocumentError
from docaudit.models import TerminologyRule
from docaudit.parser import MAX_UPLOAD_BYTES
from docaudit.readability import calculate_readability, long_sentence_findings
from docaudit.terminology import terminology_findings

@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def deny(*args, **kwargs):
        raise AssertionError("network call attempted")
    monkeypatch.setattr(socket.socket, "connect", deny)
    monkeypatch.setattr(socket, "create_connection", deny)
    import nltk
    monkeypatch.setattr(nltk, "download", deny)

@pytest.mark.parametrize("filename", ["test.txt", "test.md"])
def test_text_parser(filename):
    document = parse_document("First line.\n\nCafé works.".encode(), filename)
    assert len(document.sections) == 2
    assert document.sections[1].location.line == 3
    assert document.sections[1].text == "Café works."

def test_markdown_sections():
    document = parse_document(b"# Setup\nRun it.\n## Output\nCheck it.", "a.md")
    assert [s.title for s in document.sections] == ["Setup", "Setup", "Output", "Output"]

@pytest.mark.parametrize("data,name", [(b"", "a.txt"), (b"bad", "a.pdf"), (b"\xff", "a.txt"), (b"\x00", "a.txt"), (b"broken", "a.docx")])
def test_invalid_inputs(data, name):
    with pytest.raises(DocumentError):
        parse_document(data, name)

def test_size_limit():
    with pytest.raises(DocumentError, match="invalid_upload_size"):
        parse_document(b"a" * (MAX_UPLOAD_BYTES + 1), "a.txt")

def make_docx():
    word = WordDocument()
    word.add_heading("Setup", 1)
    word.add_paragraph("Run it.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Input"
    table.cell(0, 1).text = "Output"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    word.add_paragraph("Finish.")
    output = io.BytesIO()
    word.save(output)
    return output.getvalue()

def test_docx_order_tables_and_locations():
    document = parse_document(make_docx(), "a.docx")
    assert [s.text for s in document.sections] == ["Setup", "Run it.", "Input", "Nested", "Output", "Finish."]
    assert document.sections[2].location.table == 1
    assert document.sections[4].location.column == 2
    assert document.sections[-1].title == "Setup"

@pytest.mark.parametrize("name,content", [("word/vbaProject.bin", b"bad"), ("word/embeddings/object.bin", b"bad"), ("evil.xml", b'<!DOCTYPE x><x/>'), ("utf16.xml", '<!DOCTYPE x><x/>'.encode("utf-16")), ("../evil", b"bad")])
def test_unsafe_docx(name, content):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(make_docx())) as source, zipfile.ZipFile(output, "w") as dest:
        for entry in source.infolist():
            dest.writestr(entry, source.read(entry))
        dest.writestr(name, content)
    with pytest.raises(DocumentError):
        parse_document(output.getvalue(), "a.docx")

def test_terminology_boundaries_phrases_and_offsets():
    document = parse_document(b"Utilize utilization in   order to act. IN ORDER TO act.", "a.txt")
    config = AuditConfig(terminology=(TerminologyRule(term="utilize", replacement="use"), TerminologyRule(term="in order to", replacement="to")))
    findings = terminology_findings(document, config)
    assert len(findings) == 3
    assert findings[0].start == 0 and findings[0].end == 7
    assert all(f.location.line == 1 for f in findings)

def test_case_sensitive_and_literal_terms():
    document = parse_document(b"API api C++ C+", "a.txt")
    config = AuditConfig(terminology=(TerminologyRule(term="API", replacement="interface", case_sensitive=True), TerminologyRule(term="C++", replacement="language")))
    assert len(terminology_findings(document, config)) == 2

def test_readability_and_sentence_limit():
    document = parse_document(b"Run the tool. Check the output now.", "a.txt")
    metrics = calculate_readability(document)
    assert metrics.word_count == 7
    assert metrics.flesch_reading_ease == round(206.835 - 1.015 * 7 / 2 - 84.6 * 8 / 7, 2)
    assert metrics.flesch_kincaid_grade == round(0.39 * 7 / 2 + 11.8 * 8 / 7 - 15.59, 2)
    findings = long_sentence_findings(document, 3)
    assert len(findings) == 1
    assert findings[0].start == 14
    assert not long_sentence_findings(document, 4)

def test_empty_readability():
    metrics = calculate_readability(parse_document(b" \n", "a.txt"))
    assert metrics.word_count == 0 and metrics.flesch_reading_ease is None

def test_engine_deterministic_json():
    document = parse_document(b"Utilize the tool in order to run it.", "a.txt")
    config = load_config("config/style_guide.yaml")
    result = audit(document, config)
    assert result == audit(document, config)
    assert result.model_validate_json(result.model_dump_json()) == result
    assert len(result.findings) == 2
    assert result.section_count == 1

def test_configuration_validation():
    with pytest.raises(ValidationError):
        AuditConfig(max_sentence_words=0)
    with pytest.raises(ValidationError):
        AuditConfig(unknown=True)
    with pytest.raises(TypeError):
        audit(parse_document(b"Hello.", "a.txt"), {})

def test_engine_has_no_streamlit_imports():
    import ast
    from pathlib import Path
    for file in Path("src/docaudit").glob("*.py"):
        tree = ast.parse(file.read_text())
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                assert all(not alias.name.startswith("streamlit") for alias in node.names)
            if isinstance(node, ast.ImportFrom):
                assert not (node.module or "").startswith("streamlit")


def test_docx_expansion_limit(monkeypatch):
    import docaudit.parser as parser
    monkeypatch.setattr(parser, "MAX_EXPANDED_BYTES", 10)
    with pytest.raises(DocumentError, match="expanded_size_limit"):
        parse_document(make_docx(), "a.docx")

def test_yaml_is_safe(tmp_path):
    path = tmp_path / "rules.yaml"
    path.write_text("!!python/object/apply:os.system ['echo unsafe']")
    import yaml
    with pytest.raises(yaml.YAMLError):
        load_config(path)

def test_blank_terminology_rejected():
    with pytest.raises(ValidationError):
        TerminologyRule(term="   ", replacement="use")


def test_markdown_masks_code_and_links_preserving_offsets():
    raw = "# **Setup**\nUtilize [the tool](https://example.org/utilize).\n\n```python\nutilize()\n```\n\n    utilize()\n\nUse `utilize` instead."
    document = parse_document(raw.encode(), "a.md")
    assert document.sections[0].title == "Setup"
    assert document.sections[0].kind == "heading"
    findings = terminology_findings(document, AuditConfig(terminology=(TerminologyRule(term="utilize", replacement="use"),)))
    assert len(findings) == 1
    assert findings[0].location.line == 2
    assert document.sections[findings[0].section_index].text[findings[0].start:findings[0].end] == "Utilize"
    assert "https" not in " ".join(s.text for s in document.sections)


def test_markdown_wrapped_paragraph_and_multiline_location():
    document = parse_document(b"Setup\n=====\nRun the tool\nin order to finish.\n\n- Utilize it.\n\n| Name | Value |\n| --- | --- |\n| Mode | Safe |", "a.md")
    assert document.sections[0].title == "Setup"
    assert document.sections[1].text == "Run the tool\nin order to finish."
    assert [s.kind for s in document.sections] == ["heading", "paragraph", "list", "table", "table"]
    config = AuditConfig(terminology=(TerminologyRule(term="in order to", replacement="to"),))
    assert terminology_findings(document, config)[0].location.line == 4
    assert len(long_sentence_findings(document, 5)) == 1


@pytest.mark.parametrize("text,expected", [
    ("Use version 1.2.3. Then run it.", ["Use version 1.2.3.", "Then run it."]),
    ("Dr. Smith uses e.g. examples. Check them.", ["Dr. Smith uses e.g. examples.", "Check them."]),
    ('Say "Run it." Then stop!', ['Say "Run it."', "Then stop!"]),
    ("Visit https://example.org/docs. Then stop.", ["Visit https://example.org/docs.", "Then stop."]),
    ("One\nwrapped sentence. Final fragment", ["One\nwrapped sentence.", "Final fragment"]),
    ("   ", []),
])
def test_sentence_segmentation(text, expected):
    from docaudit.sentences import sentence_spans
    spans = list(sentence_spans(text))
    assert [text[a:b] for a, b in spans] == expected


def test_readability_uses_same_sentence_boundaries():
    document = parse_document(b"Dr. Smith runs the tool. Stop.", "a.txt")
    assert len(long_sentence_findings(document, 4)) == 1
    metrics = calculate_readability(document)
    from pyphen import Pyphen
    from docaudit.readability import WORD
    text = document.sections[0].text
    syllables = sum(len(Pyphen(lang="en_GB").positions(w)) + 1 for w in WORD.findall(text))
    assert metrics.flesch_reading_ease == round(206.835 - 1.015 * metrics.word_count / 2 - 84.6 * syllables / metrics.word_count, 2)


def test_report_filtering_and_metrics():
    from streamlit.testing.v1 import AppTest
    from pathlib import Path
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    document = parse_document(b"Utilize the tool. This sentence contains more than five words.", "a.txt")
    config = AuditConfig(max_sentence_words=5, terminology=(TerminologyRule(term="utilize", replacement="use"),))
    app.session_state["audit_result"] = audit(document, config)
    app.session_state["audit_document"] = document
    app.run()
    assert not app.exception
    assert len(app.metric) == 4
    assert len(app.dataframe[0].value) == 2
    app.selectbox[0].select("Terminology").run()
    assert not app.exception
    assert len(app.dataframe[0].value) == 1
    assert app.dataframe[0].value.iloc[0]["Matched text"] == "Utilize"
    app.selectbox[0].select("Long sentences").run()
    assert len(app.dataframe[0].value) == 1


@pytest.mark.parametrize("count,expected", [(26, "minor"), (30, "minor"), (31, "review"), (50, "review"), (51, "high")])
def test_sentence_priorities(count, expected):
    from docaudit.reporting import sentence_priority
    assert sentence_priority(count, 25) == expected


def test_prioritized_result_metadata_and_context():
    from docaudit.reporting import prioritized_findings, finding_context
    document = parse_document((" ".join(["word"] * 26) + ". " + " ".join(["word"] * 51) + ".").encode(), "a.txt")
    result = audit(document, AuditConfig())
    assert [f.priority for f in result.findings] == ["minor", "high"]
    prioritized = prioritized_findings(result.findings)
    assert prioritized[0].word_count == 51
    assert prioritized[0].word_limit == 25
    assert result.findings[0].word_count == 26
    assert finding_context(document, prioritized[0])["matched_text"] == " ".join(["word"] * 51) + "."
    payload = result.model_dump_json()
    assert '"schema_version":"2.0"' in payload
    assert "word word" not in payload


def test_report_section_navigation_and_empty_filter():
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    document = parse_document(b"# Setup\nUtilize it.\n\n# Output\nSimply run it.", "a.md")
    app.session_state["audit_document"] = document
    app.session_state["audit_result"] = audit(document, load_config("config/style_guide.yaml"))
    app.run()
    assert not app.exception
    assert len(app.dataframe) == 2
    app.selectbox[1].select("Output").run()
    assert len(app.dataframe) == 1
    assert app.dataframe[0].value.iloc[0]["Matched text"] == "Simply"
    app.multiselect[0].set_value([]).run()
    assert not app.exception
    assert len(app.dataframe) == 0
    assert app.info


def test_academic_sentence_boundaries():
    from docaudit.sentences import sentence_spans
    text = "Use the AI RMF 1.0 in the U.S. regulatory context (Smith, 2024). The plan has three requirements: first, register the baseline; second, record deviations; and third, document remediation."
    chunks = [text[a:b] for a, b in sentence_spans(text)]
    assert chunks == ["Use the AI RMF 1.0 in the U.S. regulatory context (Smith, 2024).", "The plan has three requirements: first, register the baseline; second, record deviations; and third, document remediation."]
