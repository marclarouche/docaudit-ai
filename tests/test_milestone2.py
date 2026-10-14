import io
import socket
from pathlib import Path
import pytest
from docx import Document as WordDocument
from pydantic import ValidationError
from docaudit import audit, parse_document, load_profile, load_config
from docaudit.models import AuditConfig, AcronymConfig, PreferredTerm

@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def deny(*a, **k):
        raise AssertionError("network attempted")
    monkeypatch.setattr(socket.socket, "connect", deny)
    monkeypatch.setattr(socket, "create_connection", deny)
    import nltk
    monkeypatch.setattr(nltk, "download", deny)


def acronym_audit(text):
    return audit(parse_document(text.encode(), "a.txt"), AuditConfig(max_sentence_words=1000, acronyms=AcronymConfig(enabled=True)))


def test_profiles_and_limits():
    technical = load_profile("technical", "config")
    academic = load_profile("academic", "config")
    assert technical.max_sentence_words == 25
    assert academic.max_sentence_words == 40
    text = " ".join(["word"] * 35) + "."
    document = parse_document(text.encode(), "a.txt")
    assert len(audit(document, technical).findings) == 1
    assert not audit(document, academic).findings
    with pytest.raises(ValueError):
        load_profile("../academic", "config")


def test_exclusion_remaps_original_indices_and_metrics():
    document = parse_document(b"# Intro\nUtilize it.\n\n# 9. References\nUtilize ABC.\n\n# Closing\nUtilize it.", "a.md")
    result = audit(document, load_profile("technical", "config"))
    assert result.excluded_section_count == 2
    assert result.analyzed_section_count == 4
    assert [f.section_index for f in result.findings] == [1, 5]
    assert result.findings[-1].location.line == 8
    assert result.metrics.word_count == 6
    assert result.configuration.profile == "technical"


def test_all_excluded():
    result = audit(parse_document(b"# References\nUtilize ABC.", "a.md"), AuditConfig(excluded_headings=("references*",)))
    assert not result.findings
    assert result.metrics.word_count == 0
    assert result.analyzed_section_count == 0


def test_preferred_terms_boundaries_case_and_deduplication():
    document = parse_document(b"DocAudit AI docaudit ai DocAuditAI DocAuditAIPlus", "a.txt")
    config = AuditConfig(preferred_terms=(PreferredTerm(preferred="DocAudit AI", variants=("DocAuditAI", "docaudit ai")),))
    findings = audit(document, config).findings
    assert [document.sections[0].text[f.start:f.end] for f in findings] == ["docaudit ai", "DocAuditAI"]
    assert all(f.rule == "terminology" for f in findings)


@pytest.mark.parametrize("text", [
    "Social Impact Overlay Framework (SIOF) governs risk. SIOF works.",
    "SIOF (Social Impact Overlay Framework) governs risk.",
    "National Institute of Standards and Technology (NIST) guides work. NIST helps.",
    "SIOF: Social Impact Overlay Framework",
    "AI Risk Management Framework (AI RMF 1.0) governs risk.",
])
def test_valid_acronym_definitions(text):
    result = audit(parse_document(text.encode(), "a.txt"), AuditConfig(max_sentence_words=1000, acronyms=AcronymConfig(enabled=True, allowlist=("AI",))))
    assert not result.findings


def test_late_definition_does_not_erase_first_use():
    result = acronym_audit("SIOF works.\nSocial Impact Overlay Framework (SIOF) governs risk. SIOF works.")
    assert len(result.findings) == 1
    assert result.findings[0].start == 0
    assert result.findings[0].rule == "undefined_acronym"


def test_false_expansion_and_once_per_token():
    result = acronym_audit("Unrelated Words (SIOF) is not a definition. SIOF repeats.")
    assert len(result.findings) == 1


def test_allowed_acronyms_and_control_identifiers():
    document = parse_document(b"AI uses API and AC-1(SI), CA-7(SI), and HTTP.", "a.txt")
    result = audit(document, AuditConfig(acronyms=AcronymConfig(enabled=True, allowlist=("AI", "API", "HTTP"))))
    assert not result.findings


def test_glossary_definitions_in_docx():
    word = WordDocument()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "SIOF"
    table.cell(0, 1).text = "Social Impact Overlay Framework"
    word.add_paragraph("SIOF governs risk.")
    stream = io.BytesIO()
    word.save(stream)
    result = audit(parse_document(stream.getvalue(), "a.docx"), AuditConfig(acronyms=AcronymConfig(enabled=True)))
    assert not result.findings


def test_heading_definition_is_recognized():
    document = parse_document(b"# Social Impact Overlay Framework (SIOF)\nSIOF governs risk.", "a.md")
    assert not audit(document, AuditConfig(acronyms=AcronymConfig(enabled=True))).findings


def test_excluded_definition_does_not_hide_use():
    document = parse_document(b"# References\nSocial Impact Overlay Framework (SIOF)\n\n# Body\nSIOF works.", "a.md")
    result = audit(document, load_profile("academic", "config"))
    assert len(result.findings) == 1
    assert result.findings[0].section_index == 3


@pytest.mark.parametrize("kwargs", [{"excluded_headings": (" ",)}, {"preferred_terms": [{"preferred":" "}]}, {"acronyms": {"enabled":True, "allowlist":["lowercase"]}}])
def test_invalid_configuration(kwargs):
    with pytest.raises(ValidationError):
        AuditConfig(**kwargs)


def test_profile_switch_clears_stale_report():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    document = parse_document(b"Utilize ABC.", "a.txt")
    app.session_state["audit_document"] = document
    app.session_state["audit_result"] = audit(document, load_profile("technical", "config"))
    app.run()
    app.selectbox[0].select("Undefined acronyms").run()
    assert len(app.dataframe[0].value) == 1
    app.radio[0].set_value("Academic writing").run()
    assert not app.exception
    assert not app.metric


def test_uppercase_function_labels_allowed_in_profiles():
    document = parse_document(b"GOVERN sets policy. MAP sets context. MEASURE provides evidence. MANAGE responds.", "a.txt")
    assert not audit(document, load_profile("technical", "config")).findings


def test_acronym_expansion_is_local_not_arbitrary_parentheses():
    assert len(acronym_audit("Risk Management Framework provides context (a note about RMF). RMF works.").findings) == 1


def test_excluded_heading_includes_subheadings():
    document = parse_document(b"# References\n## Sources\nUtilize XYZ.\n# Conclusion\nUtilize it.", "a.md")
    result = audit(document, load_profile("academic", "config"))
    assert result.excluded_section_count == 3
    assert len(result.findings) == 1
    assert result.findings[0].section_index == 4


def test_docx_excluded_heading_includes_subheadings():
    word = WordDocument()
    word.add_heading("References", 1)
    word.add_heading("Sources", 2)
    word.add_paragraph("Utilize XYZ.")
    word.add_heading("Conclusion", 1)
    word.add_paragraph("Utilize it.")
    stream = io.BytesIO()
    word.save(stream)
    result = audit(parse_document(stream.getvalue(), "a.docx"), load_profile("academic", "config"))
    assert result.excluded_section_count == 3
    assert len(result.findings) == 1


@pytest.mark.parametrize("profile", ["technical", "academic"])
def test_docx_samples(profile):
    config = load_profile(profile, "config")
    clean = Path("samples/milestone2-clean.docx")
    problematic = Path("samples/milestone2-problematic.docx")
    assert not audit(parse_document(clean.read_bytes(), clean.name), config).findings
    result = audit(parse_document(problematic.read_bytes(), problematic.name), config)
    assert len(result.findings) == 5
    assert {f.rule for f in result.findings} == {"terminology", "long_sentence", "undefined_acronym"}


def test_legacy_export_is_readable():
    from docaudit.models import AuditResult
    old = {"schema_version": "1.0", "document_format": "txt", "section_count": 1,
           "metrics": {"word_count": 1, "flesch_reading_ease": 100.0, "flesch_kincaid_grade": 0.0}, "findings": []}
    assert AuditResult.model_validate(old).schema_version == "1.0"


def test_profile_audit_is_deterministic_and_roundtrips():
    from docaudit.models import AuditResult
    config = load_profile("technical", "config")
    document = parse_document(b"Use DocAuditAI with XYZ.", "a.txt")
    result = audit(document, config)
    assert result == audit(document, config)
    assert AuditResult.model_validate_json(result.model_dump_json()) == result
