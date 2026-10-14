"""Bounded parsing of untrusted uploads; no files are extracted or executed."""
import io
import re
import zipfile
from pathlib import Path
from docx import Document as WordDocument
from docx.table import Table
from .models import Document, Section, Location
from .markdown import parse_markdown as _parse_markdown

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 40 * 1024 * 1024

class DocumentError(ValueError):
    """Safe error with a diagnostic code, never imported text."""

def parse_document(data: bytes, filename: str) -> Document:
    if not isinstance(data, bytes) or not data or len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("invalid_upload_size")
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in {"txt", "md", "docx"}:
        raise DocumentError("unsupported_format")
    if extension == "docx":
        return _parse_docx(data)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DocumentError("invalid_utf8") from None
    if any(ord(c) < 32 and c not in "\n\r\t" for c in text):
        raise DocumentError("binary_text")
    if extension == "md":
        return _parse_markdown(text)
    title = "Document"
    sections = []
    for number, line in enumerate(text.splitlines(), 1):
        heading = re.match(r"^ {0,3}#{1,6}\s+(.+)", line) if extension == "md" else None
        if heading:
            title = heading.group(1).strip()
        if line.strip():
            sections.append(Section(title=title, text=line, location=Location(line=number)))
    return Document(format=extension, sections=tuple(sections))

def _parse_docx(data: bytes) -> Document:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [item.filename for item in entries]
            if len(entries) > 2000 or len(set(names)) != len(names):
                raise DocumentError("invalid_archive_entries")
            if sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise DocumentError("expanded_size_limit")
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise DocumentError("invalid_docx_package")
            for item in entries:
                name = item.filename.lower()
                if item.flag_bits & 1 or name.startswith("/") or ".." in name.split("/"):
                    raise DocumentError("unsafe_archive")
                if "vbaproject" in name or "/embeddings/" in name or "/activex/" in name:
                    raise DocumentError("active_content")
                if name.endswith(".xml") or name.endswith(".rels"):
                    xml = archive.read(item)
                    if b"\x00" in xml or b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise DocumentError("unsafe_xml")
                    if name == "[content_types].xml" and b"macroenabled" in xml.lower():
                        raise DocumentError("macro_package")
        word = WordDocument(io.BytesIO(data))
        sections = []
        title = "Document"
        heading_stack = []
        paragraph = table_number = 0
        def visit(container, table_path=None, depth=0):
            nonlocal paragraph, table_number, title
            if depth > 20:
                raise DocumentError("table_nesting_limit")
            for block in container.iter_inner_content():
                if isinstance(block, Table):
                    table_number += 1
                    current_table = table_number
                    seen = set()
                    for r, row in enumerate(block.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            if cell._tc in seen:
                                continue
                            seen.add(cell._tc)
                            visit(cell, (current_table, r, c), depth + 1)
                else:
                    paragraph += 1
                    if not table_path and block.style and block.style.name.startswith("Heading"):
                        title = block.text or title
                        heading_number = re.search(r"\d+$", block.style.name)
                        level = int(heading_number.group()) if heading_number else 1
                        while heading_stack and heading_stack[-1][0] >= level:
                            heading_stack.pop()
                        heading_stack.append((level, title))
                    if block.text.strip():
                        location = Location(paragraph=paragraph)
                        if table_path:
                            location = Location(paragraph=paragraph, table=table_path[0], row=table_path[1], column=table_path[2])
                        sections.append(Section(kind="heading" if not table_path and block.style and block.style.name.startswith("Heading") else "table" if table_path else "paragraph", heading_path=tuple(t for _, t in heading_stack), title=title, text=block.text, location=location))
        visit(word)
        return Document(format="docx", sections=tuple(sections))
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("invalid_docx") from None
